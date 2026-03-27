"""Check if all Function names from the Excel findings exist in FunctionLookup.csv."""
import csv as csv_mod
import openpyxl

# 1. Load function names from Excel
xlsx = r"_analysis\x02-updateLineNums\Findings_WithCodeLine_WithTriageComment_Perfect.xlsx"
wb = openpyxl.load_workbook(xlsx, read_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
func_idx = headers.index("Function")
funcs = set()
for row in ws.iter_rows(min_row=2):
    val = row[func_idx].value
    if val:
        funcs.add(str(val).strip())
wb.close()

# 2. Load FunctionLookup.csv
csv_path = r"C:\code\codeQL_CoD\codeql\FunctionLookup.csv"
bare_names = set()
qualified_names = set()
with open(csv_path, "r", encoding="utf-8") as f:
    reader = csv_mod.reader(f)
    next(reader)  # skip header
    for row in reader:
        if len(row) >= 2:
            qualified_names.add(row[0].strip())
            bare_names.add(row[1].strip())

print(f"Excel: {len(funcs)} unique functions")
print(f"FunctionLookup.csv: {len(bare_names):,} unique bare names, {len(qualified_names):,} qualified names")
print()

# 3. Check each
found = []
missing = []
for name in sorted(funcs):
    if name in bare_names or name in qualified_names:
        found.append(name)
    else:
        bare = name.split("::")[-1]
        if bare in bare_names:
            found.append(name)
        else:
            missing.append(name)

print(f"FOUND: {len(found)}/{len(funcs)} ({100*len(found)/len(funcs):.1f}%)")
print(f"MISSING: {len(missing)}/{len(funcs)} ({100*len(missing)/len(funcs):.1f}%)")
print()
if missing:
    print("Missing functions:")
    for m in missing:
        print(f"  {m}")
