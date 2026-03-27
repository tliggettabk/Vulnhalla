"""
Compare orchestrated engine verdicts against human researcher ground truth.

Reads:
  - Human verdicts from the Excel file (CID, Report, Last Triage Comment)
  - Engine verdicts from run_report.md files in output/results_orchestrated/c/{CID}/

Usage:
    python _analysis/tools/compare_orchestrated_vs_human.py

Optional args:
    --excel PATH   Path to the Excel ground truth file (default: standard location)
    --results DIR  Path to orchestrated results root (default: output/results_orchestrated/c)
    --run RUN      Run folder name to compare (default: latest run_NNN per CID)
"""

import argparse
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


# --- Status code mapping ---
STATUS_MAP = {
    "1337": "TP (True Positive)",
    "1007": "FP (False Positive)",
    "7331": "Need More Data",
    "7337": "Conflicting Evidence",
}


def parse_args():
    parser = argparse.ArgumentParser(description="Compare orchestrated verdicts vs human ground truth")
    parser.add_argument(
        "--excel",
        default=r"_analysis\x02-updateLineNums\Findings_WithCodeLine_WithTriageComment_Perfect.xlsx",
        help="Path to Excel file with human verdicts",
    )
    parser.add_argument(
        "--results",
        default=r"output\results_orchestrated\c",
        help="Path to orchestrated results directory",
    )
    parser.add_argument(
        "--run",
        default=None,
        help="Specific run folder (e.g. run_001). Default: latest per CID",
    )
    return parser.parse_args()


def load_human_verdicts(excel_path: str) -> dict:
    """Load CID -> {report, triage_comment} from Excel."""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb.active
    headers = {cell.value: idx for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}

    cid_idx = headers["CID"]
    report_idx = headers["Report"]
    triage_idx = headers["Last Triage Comment"]

    verdicts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[cid_idx]
        if cid is None:
            continue
        cid = str(int(cid))
        verdicts[cid] = {
            "report": row[report_idx],  # "Yes" = human says TP / reportable
            "triage_comment": row[triage_idx] or "",
        }
    wb.close()
    return verdicts


def get_latest_run(cid_dir: str) -> str | None:
    """Find the latest run_NNN folder in a CID directory."""
    if not os.path.isdir(cid_dir):
        return None
    runs = sorted(
        [d for d in os.listdir(cid_dir) if d.startswith("run_") and os.path.isdir(os.path.join(cid_dir, d))]
    )
    return runs[-1] if runs else None


def extract_verdict_from_report(report_path: str) -> str | None:
    """Extract the status code (1337/1007/7331/7337) from run_report.md."""
    if not os.path.isfile(report_path):
        return None
    with open(report_path, "r", encoding="utf-8") as f:
        content = f.read()
    # Look for **1337**, **1007**, **7331**, **7337** near the end
    matches = re.findall(r"\*\*(\d{4})\*\*", content)
    for code in reversed(matches):
        if code in STATUS_MAP:
            return code
    return None


def extract_summary_stats(run_dir: str) -> dict:
    """Pull key stats from run_summary.json."""
    summary_path = os.path.join(run_dir, "run_summary.json")
    if not os.path.isfile(summary_path):
        return {}
    with open(summary_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return {
        "tokens": data.get("total_tokens", 0),
        "cost": data.get("estimated_cost_usd", 0),
        "duration": data.get("duration_seconds", 0),
        "leads": data.get("leads_count", 0),
        "tool_calls": data.get("tool_calls", 0),
    }


def human_verdict_label(report_value: str) -> str:
    """Convert human Report column to a label."""
    if report_value and str(report_value).strip().lower() == "yes":
        return "TP (Reportable)"
    elif report_value and str(report_value).strip().lower() == "no":
        return "FP (Not Reportable)"
    return f"Unknown ({report_value})"


def compare(human_verdicts: dict, results_dir: str, run_name: str | None) -> list[dict]:
    """Compare engine verdicts against human for all CIDs with orchestrated results."""
    comparisons = []

    # Find all CIDs with orchestrated results
    if not os.path.isdir(results_dir):
        print(f"ERROR: Results directory not found: {results_dir}")
        return comparisons

    for cid in sorted(os.listdir(results_dir)):
        cid_dir = os.path.join(results_dir, cid)
        if not os.path.isdir(cid_dir):
            continue

        # Determine which run to use
        run = run_name or get_latest_run(cid_dir)
        if not run:
            continue

        run_dir = os.path.join(cid_dir, run)
        report_path = os.path.join(run_dir, "run_report.md")

        engine_code = extract_verdict_from_report(report_path)
        stats = extract_summary_stats(run_dir)

        human = human_verdicts.get(cid, {})
        human_label = human_verdict_label(human.get("report")) if human else "N/A (not in Excel)"

        # Determine agreement
        if not human:
            agreement = "N/A"
        elif engine_code is None:
            agreement = "N/A (no engine verdict)"
        elif engine_code in ("7331", "7337"):
            agreement = "INCONCLUSIVE"
        elif human.get("report", "").strip().lower() == "yes" and engine_code == "1337":
            agreement = "AGREE"
        elif human.get("report", "").strip().lower() == "yes" and engine_code == "1007":
            agreement = "DISAGREE"
        elif human.get("report", "").strip().lower() == "no" and engine_code == "1007":
            agreement = "AGREE"
        elif human.get("report", "").strip().lower() == "no" and engine_code == "1337":
            agreement = "DISAGREE"
        else:
            agreement = "UNKNOWN"

        comparisons.append({
            "cid": cid,
            "run": run,
            "human_label": human_label,
            "human_comment": human.get("triage_comment", ""),
            "engine_code": engine_code,
            "engine_label": STATUS_MAP.get(engine_code, "N/A") if engine_code else "N/A",
            "agreement": agreement,
            **stats,
        })

    return comparisons


def print_report(comparisons: list[dict]):
    """Print comparison report to stdout."""
    if not comparisons:
        print("No comparisons to report.")
        return

    # Summary table
    print("=" * 100)
    print("ORCHESTRATED ENGINE vs HUMAN RESEARCHER — Comparison Report")
    print("=" * 100)
    print()

    # Header
    print(f"{'CID':<8} {'Run':<10} {'Human':<20} {'Engine':<22} {'Match':<14} {'Cost':>7} {'Tokens':>8} {'Leads':>5}")
    print("-" * 100)

    agree = disagree = inconclusive = 0
    total_cost = 0.0
    total_tokens = 0

    for c in comparisons:
        cost_str = f"${c.get('cost', 0):.2f}"
        tokens_str = f"{c.get('tokens', 0):,}"
        print(
            f"{c['cid']:<8} {c['run']:<10} {c['human_label']:<20} {c['engine_label']:<22} "
            f"{c['agreement']:<14} {cost_str:>7} {tokens_str:>8} {c.get('leads', ''):>5}"
        )
        if c["agreement"] == "AGREE":
            agree += 1
        elif c["agreement"] == "DISAGREE":
            disagree += 1
        elif c["agreement"] == "INCONCLUSIVE":
            inconclusive += 1
        total_cost += c.get("cost", 0)
        total_tokens += c.get("tokens", 0)

    print("-" * 100)
    total = len(comparisons)
    print(f"\nResults: {agree} agree, {disagree} disagree, {inconclusive} inconclusive out of {total} CIDs")
    if agree + disagree > 0:
        accuracy = agree / (agree + disagree) * 100
        print(f"Accuracy (excluding inconclusive): {accuracy:.0f}% ({agree}/{agree + disagree})")
    print(f"Total cost: ${total_cost:.2f} | Total tokens: {total_tokens:,}")

    # Detail section
    print("\n" + "=" * 100)
    print("DETAILED COMPARISON")
    print("=" * 100)

    for c in comparisons:
        icon = {"AGREE": "✓", "DISAGREE": "✗", "INCONCLUSIVE": "?"}.get(c["agreement"], "—")
        print(f"\n{icon} CID {c['cid']} ({c['run']})")
        print(f"  Human:  {c['human_label']}")
        print(f"  Engine: {c['engine_label']}")
        print(f"  Match:  {c['agreement']}")
        if c["human_comment"]:
            comment = c["human_comment"][:200]
            print(f"  Human rationale: {comment}")
        print(f"  Stats: {c.get('leads', '?')} leads, {c.get('tokens', 0):,} tokens, ${c.get('cost', 0):.2f}, {c.get('duration', 0):.0f}s")


def main():
    args = parse_args()
    print(f"Loading human verdicts from: {args.excel}")
    human = load_human_verdicts(args.excel)
    print(f"  Found {len(human)} CIDs in Excel\n")

    comparisons = compare(human, args.results, args.run)
    print_report(comparisons)


if __name__ == "__main__":
    main()
