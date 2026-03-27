"""Append Development Journey slides to the existing PPTX.

Usage:  python _rearch/insert_devjourney_slides.py

Appends slides at the end. Backs up first. Does NOT touch existing slides.
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-journey.bak.pptx"

# ── Colours ──
BG        = RGBColor(0x1E, 0x1E, 0x2E)
TITLE_CLR = RGBColor(0x89, 0xB4, 0xFA)
TEXT_CLR  = RGBColor(0xCD, 0xD6, 0xF4)
ACCENT    = RGBColor(0xA6, 0xE3, 0xA1)
WARN      = RGBColor(0xFA, 0xB3, 0x87)
CODE_BG   = RGBColor(0x31, 0x32, 0x44)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DIM       = RGBColor(0x6C, 0x70, 0x86)
MAUVE     = RGBColor(0xCB, 0xA6, 0xF7)
YELLOW    = RGBColor(0xF9, 0xE2, 0xAF)
TEAL      = RGBColor(0x94, 0xE2, 0xD5)
RED_CLR   = RGBColor(0xF3, 0x8B, 0xA8)

# ── Helpers ──

def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, text, size=18, color=TEXT_CLR,
             bold=False, align=PP_ALIGN.LEFT, font_name="Cascadia Code"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


def add_para(tf, text, size=18, color=TEXT_CLR, bold=False,
             font_name="Cascadia Code", space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.space_before = space_before
    return p


def add_title(slide, text):
    add_text(slide, 0.5, 0.3, 12, 0.7, text, size=32,
             color=TITLE_CLR, bold=True, font_name="Segoe UI")


def add_subtitle(slide, text, top=1.1):
    add_text(slide, 0.5, top, 12, 0.5, text, size=18,
             color=TEXT_CLR, font_name="Segoe UI")


def add_rounded_box(slide, left, top, width, height, fill_color, text,
                    text_size=16, text_color=WHITE, bold=True,
                    font_name="Segoe UI", align=PP_ALIGN.CENTER):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(text_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


# ══════════════════════════════════════════════════════════════════════
# Load & backup
# ══════════════════════════════════════════════════════════════════════
shutil.copy2(TARGET, BACKUP)
print(f"Backed up \u2192 {BACKUP.name}")

prs = Presentation(str(TARGET))
n_before = len(prs.slides)
print(f"Existing slides: {n_before}")

# ══════════════════════════════════════════════════════════════════════
# SLIDE A — Development Journey (timeline)
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Development Journey \u2014 Before Orchestrator")
add_subtitle(s, "The path from first contact with Vulnhalla to identifying the need for a new architecture")

# Phase boxes (vertical timeline, left side)
phases = [
    (MAUVE, "1",
     "Get to Know Vulnhalla",
     "Run on stub C++ programs with known flaws. "
     "Fix Windows-specific bugs. Add logging for transparency."),
    (TEAL, "2",
     "First Run on Coverity CoD Results",
     "Sync findings line numbers with 10/2025 CodeQL DB. "
     "Generate lookup CSVs. Fix FunctionTree.csv generation. "
     "Trim findings to those found in CSV. Add LLM usage guards."),
    (WARN, "3",
     "Enhance Vulnhalla",
     "Create regression testing framework. Tune prompts: "
     "ignore assertions, prevent duplicate tool use, require "
     "tool-call justification, enforce known-facts grounding."),
    (RED_CLR, "4",
     "Fatal Flaw \u2014 Need State Management",
     "Single-conversation mode couldn\u2019t manage evidence "
     "reliably. LLM drifted, repeated tools, lost track of "
     "established facts. Prompt tuning hit its ceiling."),
    (ACCENT, "5",
     "Add Orchestrator Mode",
     "Decompose into Plan \u2192 Investigate \u2192 Synthesize. "
     "Python controls the process. LLM answers focused questions."),
]

for i, (color, num, title, desc) in enumerate(phases):
    top = 1.55 + i * 1.1
    # Number circle
    add_rounded_box(s, 0.6, top, 0.5, 0.5, color, num,
                    text_size=18, text_color=WHITE, bold=True)
    # Title
    add_text(s, 1.3, top - 0.05, 4, 0.5, title,
             size=18, color=color, bold=True, font_name="Segoe UI")
    # Description
    add_text(s, 1.3, top + 0.35, 4.5, 0.7, desc,
             size=12, color=TEXT_CLR, font_name="Segoe UI")

# Right side: Key lessons learned
tf_right = add_text(s, 7, 1.5, 5.5, 0.5,
                    "Key Lessons Along the Way", size=22,
                    color=YELLOW, bold=True, font_name="Segoe UI")

lessons = [
    ("Assertions fooled the LLM",
     "LLMs treat assert() as a real runtime guard. Without explicit "
     "instruction to ignore them, every assert looked like a valid bound "
     "check \u2014 leading to false FP verdicts."),
    ("Duplicate tool calls wasted budget",
     "The LLM frequently re-requested the same function it had already "
     "looked up rounds earlier. No built-in memory across rounds."),
    ("Prompt tuning has a ceiling",
     "Adding rules to the system prompt helped (ignore asserts, justify "
     "tool use, cite evidence) but couldn\u2019t fix the structural problem: "
     "one long conversation with no decomposition."),
    ("State management was the blocker",
     "Evidence tracking via sticky flags (MECHANISM, GUARDS, EXPLOITABILITY) "
     "was bolted on. The LLM still lost context in long conversations and "
     "contradicted its own earlier findings."),
    ("Orchestration solved the root cause",
     "Splitting into 3 phases gave Python control over state, tool budgets, "
     "and evidence flow. The LLM\u2019s job shrank from \u201cdrive the investigation\u201d "
     "to \u201canswer one question.\u201d"),
]
for title, desc in lessons:
    add_para(tf_right, f"\u2022  {title}", size=15, color=ACCENT, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf_right, f"   {desc}", size=12, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# ══════════════════════════════════════════════════════════════════════
# SLIDE B — Prompt Tuning: What Worked vs What Didn't
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Old Mode Prompt Tuning \u2014 What Worked & What Didn\u2019t")
add_subtitle(s, "Iterative prompt improvements that carried forward into the orchestrator prompts")

# Left: What worked (carried forward)
tf_left = add_text(s, 0.6, 1.5, 5.8, 0.5,
                   "\u2705  Carried Forward", size=22,
                   color=ACCENT, bold=True, font_name="Segoe UI")
worked = [
    ("Assertions are NO-OPs",
     "Explicitly telling the LLM that assert/core_assert compile to "
     "nothing in production. Applies transitively: if a guard\u2019s limit "
     "is only bounded by an assert, the guard itself is broken."),
    ("Justify every tool call",
     "Requiring a reason parameter before each tool request. Cuts "
     "speculative lookups. Became the \u201cCan I answer now?\u201d checklist "
     "in the investigator prompt."),
    ("Only look up visible names",
     "LLMs invent function names based on conventions. Forcing tool "
     "args to appear in already-returned code eliminated hallucinated "
     "lookups."),
    ("Don\u2019t infer from names",
     "GetSize() doesn\u2019t necessarily return a size. Ask about observable "
     "properties (return type, value range) instead of assuming."),
]
for title, desc in worked:
    add_para(tf_left, f"  {title}", size=16, color=ACCENT, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf_left, f"   {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# Right: What didn't work (structural limits)
tf_right = add_text(s, 7, 1.5, 5.8, 0.5,
                    "\u274c  Hit the Ceiling", size=22,
                    color=WARN, bold=True, font_name="Segoe UI")
didnt = [
    ("Evidence tracking via system messages",
     "Injecting ESTABLISHED_EVIDENCE as a system message each round "
     "worked at first, but the LLM still contradicted prior findings "
     "in longer conversations. Evidence wasn\u2019t \u201csticky\u201d enough."),
    ("Dead-end handling instructions",
     "Telling the LLM to \u201cexhaust all leads before concluding\u201d didn\u2019t "
     "prevent it from fixating on one path for many rounds while "
     "ignoring alternatives."),
    ("Sufficiency flags",
     "Tracking MECHANISM/GUARDS/EXPLOITABILITY as booleans forced "
     "premature \u201cyes\u201d answers. Once set, they couldn\u2019t be revised "
     "when new contradicting evidence appeared."),
    ("Single-conversation control flow",
     "No amount of prompt tuning could fix the core issue: the LLM "
     "decided what to investigate and in what order. It needed "
     "structure imposed from outside."),
]
for title, desc in didnt:
    add_para(tf_right, f"  {title}", size=16, color=WARN, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf_right, f"   {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# Bottom
add_text(s, 0.6, 6.3, 12, 0.5,
         "The lessons that worked became the foundation of the orchestrator prompts. "
         "The ones that didn\u2019t motivated the 3-phase architecture.",
         size=16, color=YELLOW, bold=True, font_name="Segoe UI")

# ══════════════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════════════
prs.save(str(TARGET))
n_after = len(prs.slides)
print(f"Added {n_after - n_before} slides ({n_before} \u2192 {n_after})")
print(f"Saved \u2192 {TARGET.name}")
