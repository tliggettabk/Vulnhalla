"""Append CSV & Tool Enhancement slides to the existing PPTX.

Usage:  python _rearch/insert_csv_slides.py

This script:
1. Loads the existing Vulnhalla_Orchestrated_Overview.pptx
2. Appends new slides at the end (does NOT touch existing slides)
3. Saves in-place (backs up to .bak.pptx first)
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-csv.bak.pptx"

# ── Colours (match existing deck) ──
BG        = RGBColor(0x1E, 0x1E, 0x2E)
TITLE_CLR = RGBColor(0x89, 0xB4, 0xFA)
TEXT_CLR  = RGBColor(0xCD, 0xD6, 0xF4)
ACCENT    = RGBColor(0xA6, 0xE3, 0xA1)
WARN      = RGBColor(0xFA, 0xB3, 0x87)
CODE_BG   = RGBColor(0x31, 0x32, 0x44)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DIM       = RGBColor(0x6C, 0x70, 0x86)
TBL_HDR   = RGBColor(0x45, 0x47, 0x5A)
TBL_ROW1  = RGBColor(0x2A, 0x2B, 0x3D)
TBL_ROW2  = RGBColor(0x24, 0x25, 0x35)
MAUVE     = RGBColor(0xCB, 0xA6, 0xF7)
YELLOW    = RGBColor(0xF9, 0xE2, 0xAF)
TEAL      = RGBColor(0x94, 0xE2, 0xD5)

# ── Helpers (same as main builder) ──

def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, text, size=18, color=TEXT_CLR,
             bold=False, align=PP_ALIGN.LEFT, font_name="Cascadia Code"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


def add_para(tf, text, size=18, color=TEXT_CLR, bold=False,
             font_name="Cascadia Code", space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.space_before = space_before
    return p


def add_title(slide, text):
    add_text(slide, 0.5, 0.3, 12, 0.7, text, size=32,
             color=TITLE_CLR, bold=True, font_name="Segoe UI")


def add_subtitle(slide, text, top=1.1):
    add_text(slide, 0.5, top, 12, 0.5, text, size=18,
             color=TEXT_CLR, font_name="Segoe UI")


def styled_table(slide, rows, cols, left, top, width, height,
                 col_widths=None):
    tbl_shape = slide.shapes.add_table(rows, cols,
                                       Inches(left), Inches(top),
                                       Inches(width), Inches(height))
    tbl = tbl_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Inches(w)
    for row_idx, row in enumerate(tbl.rows):
        for cell in row.cells:
            cell.fill.solid()
            if row_idx == 0:
                cell.fill.fore_color.rgb = TBL_HDR
            elif row_idx % 2 == 1:
                cell.fill.fore_color.rgb = TBL_ROW1
            else:
                cell.fill.fore_color.rgb = TBL_ROW2
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(14)
                p.font.color.rgb = TEXT_CLR
                p.font.name = "Cascadia Code"
    return tbl


def set_cell(tbl, r, c, text, color=TEXT_CLR, bold=False, size=14):
    cell = tbl.cell(r, c)
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Cascadia Code"


def add_rounded_box(slide, left, top, width, height, fill_color, text,
                    text_size=16, text_color=WHITE, bold=True,
                    font_name="Segoe UI", align=PP_ALIGN.CENTER):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(text_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


# ══════════════════════════════════════════════════════════════════════
# Load existing deck
# ══════════════════════════════════════════════════════════════════════
shutil.copy2(TARGET, BACKUP)
print(f"Backed up \u2192 {BACKUP.name}")

prs = Presentation(str(TARGET))
n_before = len(prs.slides)
print(f"Existing slides: {n_before}")

# ══════════════════════════════════════════════════════════════════════
# SLIDE A \u2014 Data Layer: The Problem
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Data Layer \u2014 Closing Coverage Gaps")
add_subtitle(s, "CodeQL\u2019s type model has blind spots that cause the LLM to waste tool calls on missing lookups")

# Left: The 3 problems
tf_left = add_text(s, 0.6, 1.5, 5.8, 0.5,
                   "What Was Missing", size=22,
                   color=WARN, bold=True, font_name="Segoe UI")

problems = [
    ("Enums invisible",
     "C++11 scoped enums (enum class) aren\u2019t covered by CodeQL\u2019s "
     "NameQualifyingElement. get_class() returned \u201cnot found\u201d for "
     "every enum, forcing the LLM to burn tool calls searching."),
    ("Typedefs invisible",
     "typedef and using aliases (TypedefType) also fall outside "
     "NameQualifyingElement. CID 15518 wasted 15 tool calls (60% of "
     "the budget) trying every tool to find StLayerMaskOMPV."),
    ("Macro vs. global guessing",
     "The LLM had separate get_macro and get_global_var tools but "
     "couldn\u2019t tell which one a name was \u2014 so it guessed, often "
     "wrong, burning a call before trying the other."),
]
for title, desc in problems:
    add_para(tf_left, f"\u26a0  {title}", size=17, color=WARN, bold=True,
             font_name="Segoe UI", space_before=Pt(14))
    add_para(tf_left, f"   {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(3))

# Right: Impact
tf_right = add_text(s, 7, 1.5, 5.8, 0.5,
                    "Impact on Analysis", size=22,
                    color=TITLE_CLR, bold=True, font_name="Segoe UI")
impacts = [
    ("\u2022  Wasted tool budget", "Each failed lookup costs a tool call. With a 12-call follow-up budget per lead, missing lookups can exhaust the budget before answering the question."),
    ("\u2022  Incorrect conclusions", "When the LLM can\u2019t find a type definition, it may guess what the type does based on its name \u2014 exactly what the prompts forbid."),
    ("\u2022  Cascading failures", "A missed enum or typedef often blocks downstream lookups (e.g., can\u2019t find methods on a type you can\u2019t resolve)."),
    ("\u2022  Higher cost", "More rounds, more tokens, more API calls \u2014 all for lookups that should have succeeded on the first try."),
]
for title, desc in impacts:
    add_para(tf_right, title, size=16, color=ACCENT, bold=True,
             font_name="Segoe UI", space_before=Pt(12))
    add_para(tf_right, f"   {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# ══════════════════════════════════════════════════════════════════════
# SLIDE B \u2014 New CSVs: EnumLookup & TypeAliasLookup
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "New Lookup CSVs")
add_subtitle(s, "Two new CodeQL-generated CSVs fill the coverage gaps in get_class()")

# EnumLookup box (left)
add_rounded_box(s, 0.6, 1.6, 5.8, 0.6, MAUVE,
                "EnumLookup.csv \u2014 C++11 Scoped Enums",
                text_size=18, text_color=WHITE)

tf_enum = add_text(s, 0.6, 2.4, 5.8, 0.5, "", size=15, font_name="Segoe UI")
enum_points = [
    ("Problem", "enum class types invisible to Classes.ql (not a NameQualifyingElement)"),
    ("Solution", "New CodeQL query (EnumLookup.ql) + smart structural parser"),
    ("Smart parser", "Brace-matches from source to fix CodeQL\u2019s broken end-lines. "
     "Filters forward declarations, macro enums, wrong-enum hits."),
    ("Accuracy", "100% correct code extraction (vs 78% with naive brace-matching)"),
    ("Coverage", "10,760 unique enums with correct line spans"),
]
for label, desc in enum_points:
    add_para(tf_enum, f"  {label}:  {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(6))

# TypeAliasLookup box (right)
add_rounded_box(s, 6.8, 1.6, 5.8, 0.6, TEAL,
                "TypeAliasLookup.csv \u2014 typedef / using Aliases",
                text_size=18, text_color=WHITE)

tf_alias = add_text(s, 6.8, 2.4, 5.8, 0.5, "", size=15, font_name="Segoe UI")
alias_points = [
    ("Problem", "TypedefType falls outside NameQualifyingElement \u2014 all typedefs invisible"),
    ("Solution", "New CodeQL query (TypeAliasLookup.ql) + deduplication pipeline"),
    ("CodeQL gotchas", "matches() uses SQL LIKE (_ = wildcard); @kind problem "
     "corrupts multi-column output; getLocation() silently drops rows"),
    ("Deduplication", "554K raw rows \u2192 280K (template filter) \u2192 32,303 unique"),
    ("Coverage", "Both C-style (typedef int X) and C++11 (using X = int) aliases"),
]
for label, desc in alias_points:
    add_para(tf_alias, f"  {label}:  {desc}", size=13, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(6))

# Bottom: unified search
add_text(s, 0.6, 5.8, 12, 0.8,
         "get_class() now searches all three CSVs in order:  "
         "Classes.csv \u2192 EnumLookup.csv \u2192 TypeAliasLookup.csv\n"
         "Two-pass matching: exact first (prevents substring false hits), "
         "then fuzzy if nothing found.",
         size=16, color=YELLOW, bold=True, font_name="Segoe UI")

# ══════════════════════════════════════════════════════════════════════
# SLIDE C \u2014 Tool Merge: get_macro_or_global
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Tool Merge \u2014 get_macro_or_global")
add_subtitle(s, "Two separate tools consolidated into one \u2014 the LLM no longer has to guess")

# Before / After boxes
add_rounded_box(s, 0.6, 1.6, 5.5, 0.55, WARN,
                "Before:  2 separate tools",
                text_size=18, text_color=WHITE)
tf_before = add_text(s, 0.6, 2.3, 5.5, 0.5, "", size=15, font_name="Segoe UI")
before_points = [
    "\u2022  get_macro(macro_name)  \u2014  searches Macros.csv only",
    "\u2022  get_global_var(global_var_name)  \u2014  searches GlobalVars.csv only",
    "",
    "\u26a0  LLM sees MAX_PLAYERS in code but can\u2019t tell if it\u2019s a",
    "   #define or a const int  \u2192  guesses one, fails, tries the other",
    "\u26a0  Two wasted tool calls for a single name lookup",
    "\u26a0  Misleading prompt guidance (\u201ctry get_global_var for typedefs\u201d)",
    "   sent the LLM down a dead end  \u2014 global_var never finds typedefs",
]
for line in before_points:
    clr = WARN if line.startswith("\u26a0") else TEXT_CLR
    add_para(tf_before, f"  {line}", size=14, color=clr,
             font_name="Segoe UI", space_before=Pt(5))

add_rounded_box(s, 6.8, 1.6, 5.5, 0.55, ACCENT,
                "After:  1 unified tool",
                text_size=18, text_color=WHITE)
tf_after = add_text(s, 6.8, 2.3, 5.5, 0.5, "", size=15, font_name="Segoe UI")
after_points = [
    "\u2022  get_macro_or_global(name)  \u2014  tries Macros.csv first,",
    "   falls back to GlobalVars.csv automatically",
    "",
    "\u2705  One call resolves macros and globals alike",
    "\u2705  Response prefixed [macro] or [global_var] so LLM",
    "   knows which type it found",
    "\u2705  Backward compatible \u2014 old plans using get_macro or",
    "   get_global_var still work (mapped to unified tool)",
    "\u2705  Removed misleading \u201ctry get_global_var for typedefs\u201d",
    "   guidance from prompts",
]
for line in after_points:
    clr = ACCENT if line.startswith("\u2705") else TEXT_CLR
    add_para(tf_after, f"  {line}", size=14, color=clr,
             font_name="Segoe UI", space_before=Pt(5))

# Bottom: tool count
add_text(s, 0.6, 6.2, 12, 0.5,
         "Tool count:  5 \u2192 4   (get_function_code, get_caller_function, "
         "get_class, get_macro_or_global)",
         size=16, color=YELLOW, bold=True, font_name="Segoe UI")

# ══════════════════════════════════════════════════════════════════════
# SLIDE D \u2014 Summary: Before vs After
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Data Layer \u2014 Before vs. After")

rows_data = [
    ("Aspect",              "Before",                            "After"),
    ("Enum lookups",        "Not found (invisible to get_class)", "10,760 enums in EnumLookup.csv"),
    ("Typedef lookups",     "Not found (15 calls wasted)",        "32,303 aliases in TypeAliasLookup.csv"),
    ("Macro / global",      "2 tools \u2014 LLM guesses which",  "1 unified get_macro_or_global"),
    ("get_class() sources", "Classes.csv only",                   "Classes + Enum + TypeAlias (3 CSVs)"),
    ("Matching strategy",   "Substring (false hits possible)",    "Exact-first, fuzzy fallback"),
    ("Enum end-lines",      "Wrong (CodeQL: end = start)",        "Correct (smart brace-match parser)"),
    ("Typedef dedup",       "N/A",                                "554K \u2192 32K (94% reduction)"),
    ("Prompt guidance",     "Misleading (try get_global_var)",     "Removed; typedefs in get_class()"),
]

tbl = styled_table(s, len(rows_data), 3, 0.6, 1.2, 12.1, 5.5,
                   col_widths=[2.8, 4.6, 4.7])
for r, (a, b, c) in enumerate(rows_data):
    if r == 0:
        set_cell(tbl, r, 0, a, color=TITLE_CLR, bold=True, size=14)
        set_cell(tbl, r, 1, b, color=TITLE_CLR, bold=True, size=14)
        set_cell(tbl, r, 2, c, color=TITLE_CLR, bold=True, size=14)
    else:
        set_cell(tbl, r, 0, a, color=YELLOW, bold=True, size=13)
        set_cell(tbl, r, 1, b, color=WARN, size=13)
        set_cell(tbl, r, 2, c, color=ACCENT, size=13)

# ══════════════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════════════
prs.save(str(TARGET))
n_after = len(prs.slides)
print(f"Added {n_after - n_before} slides ({n_before} \u2192 {n_after})")
print(f"Saved \u2192 {TARGET.name}")
