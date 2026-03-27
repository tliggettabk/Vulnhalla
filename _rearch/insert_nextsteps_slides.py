"""Append Next Steps / Deferred Items slides to the existing PPTX.

Usage:  python _rearch/insert_nextsteps_slides.py

Appends slides at the end. Backs up first. Does NOT touch existing slides.
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-nextsteps.bak.pptx"

# ── Colours (Catppuccin Mocha) ──
BG        = RGBColor(0x1E, 0x1E, 0x2E)
TITLE_CLR = RGBColor(0x89, 0xB4, 0xFA)
TEXT_CLR  = RGBColor(0xCD, 0xD6, 0xF4)
ACCENT    = RGBColor(0xA6, 0xE3, 0xA1)
WARN      = RGBColor(0xFA, 0xB3, 0x87)
CODE_BG   = RGBColor(0x31, 0x32, 0x44)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DIM       = RGBColor(0x6C, 0x70, 0x86)
MAUVE     = RGBColor(0xCB, 0xA6, 0xF7)
YELLOW    = RGBColor(0xF9, 0xE2, 0xAF)
TEAL      = RGBColor(0x94, 0xE2, 0xD5)
RED_CLR   = RGBColor(0xF3, 0x8B, 0xA8)

# ── Helpers ──

def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, text, size=18, color=TEXT_CLR,
             bold=False, align=PP_ALIGN.LEFT, font_name="Cascadia Code"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


def add_para(tf, text, size=18, color=TEXT_CLR, bold=False,
             font_name="Cascadia Code", space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.space_before = space_before
    return p


def add_title(slide, text):
    add_text(slide, 0.5, 0.3, 12, 0.7, text, size=32,
             color=TITLE_CLR, bold=True, font_name="Segoe UI")


def add_subtitle(slide, text, top=1.1):
    add_text(slide, 0.5, top, 12, 0.5, text, size=18,
             color=TEXT_CLR, font_name="Segoe UI")


def add_rounded_box(slide, left, top, width, height, fill_color, text,
                    text_size=16, text_color=WHITE, bold=True,
                    font_name="Segoe UI", align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(text_size)
    p.font.color.rgb = text_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return tf


def add_section_header(slide, left, top, width, icon, label, color):
    """A coloured rounded-rect section header."""
    add_rounded_box(slide, left, top, width, 0.45, color,
                    f"{icon}  {label}", text_size=18, text_color=WHITE,
                    bold=True)


# ══════════════════════════════════════════════════════════════════════
# Load & backup
# ══════════════════════════════════════════════════════════════════════
shutil.copy2(TARGET, BACKUP)
print(f"Backed up \u2192 {BACKUP.name}")

prs = Presentation(str(TARGET))
n_before = len(prs.slides)
print(f"Existing slides: {n_before}")

# ══════════════════════════════════════════════════════════════════════
# SLIDE A — Next Steps: Road to Production
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Next Steps \u2014 Road to Production")
add_subtitle(s, "From prototype success to production-scale deployment")

# --- Column 1: Immediate ---
col1_left = 0.5
add_section_header(s, col1_left, 1.65, 3.8, "\u26A1", "Immediate", ACCENT)

immediate_items = [
    ("Process full CoD finding set",
     "Expand from 19-CID test batch to the "
     "complete backlog of Coverity findings."),
    ("Tune planner for edge cases",
     "Continue prompt iteration on regressions "
     "discovered in batch runs (guard leads, "
     "data-flow prioritization)."),
    ("Validate cost at scale",
     "Confirm per-finding costs stay within budget "
     "as we move to hundreds of findings."),
]
tf1 = add_text(s, col1_left + 0.1, 2.2, 3.6, 0.3, "", size=14,
               color=TEXT_CLR, font_name="Segoe UI")
for title, desc in immediate_items:
    add_para(tf1, f"\u2022  {title}", size=14, color=ACCENT, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf1, f"   {desc}", size=12, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# --- Column 2: Near-term ---
col2_left = 4.7
add_section_header(s, col2_left, 1.65, 3.8, "\U0001F4CB", "Near-term", TEAL)

nearterm_items = [
    ("Regression test suite",
     "Automated re-run of known CIDs after prompt or "
     "code changes to catch accuracy regressions early."),
    ("Result review workflow",
     "Tooling for security engineers to review, accept, "
     "or override LLM verdicts efficiently."),
    ("Multi-language support",
     "Extend CodeQL queries and lookup CSVs to C#, "
     "Java, and other languages in the codebase."),
]
tf2 = add_text(s, col2_left + 0.1, 2.2, 3.6, 0.3, "", size=14,
               color=TEXT_CLR, font_name="Segoe UI")
for title, desc in nearterm_items:
    add_para(tf2, f"\u2022  {title}", size=14, color=TEAL, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf2, f"   {desc}", size=12, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# --- Column 3: Production ---
col3_left = 8.9
add_section_header(s, col3_left, 1.65, 3.8, "\U0001F680", "Production", MAUVE)

prod_items = [
    ("CI/CD integration",
     "Trigger orchestrated analysis automatically "
     "on new Coverity findings in the build pipeline."),
    ("Dashboard & reporting",
     "Aggregate verdicts, confidence levels, and cost "
     "metrics into a team-facing dashboard."),
    ("Feedback loop",
     "Track human overrides of LLM verdicts to continuously "
     "improve prompts and lead generation."),
]
tf3 = add_text(s, col3_left + 0.1, 2.2, 3.6, 0.3, "", size=14,
               color=TEXT_CLR, font_name="Segoe UI")
for title, desc in prod_items:
    add_para(tf3, f"\u2022  {title}", size=14, color=MAUVE, bold=True,
             font_name="Segoe UI", space_before=Pt(10))
    add_para(tf3, f"   {desc}", size=12, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))

# Bottom highlight
add_text(s, 0.5, 6.4, 12, 0.5,
         "Goal: Move from \u201ccan we do this?\u201d  \u2192  \u201crun it on everything, automatically.\u201d",
         size=17, color=YELLOW, bold=True, font_name="Segoe UI")


# ══════════════════════════════════════════════════════════════════════
# SLIDE B — Deferred for Now & Why
# ══════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_title(s, "Deferred for Now \u2014 and Why")
add_subtitle(s, "Good ideas deliberately postponed to keep the prototype focused")

# --- Left column: Performance ---
perf_left = 0.6
add_section_header(s, perf_left, 1.65, 5.6, "\u23F1", "Performance", TEAL)

perf_items = [
    ("Async agents",
     "Run investigator leads in parallel instead of sequentially. "
     "Would cut wall-clock time ~3\u00D7 for multi-lead plans.",
     "Deferred: Adds concurrency complexity (rate limits, error "
     "handling, token accounting). Sequential is fast enough for "
     "the current batch size and simpler to debug."),
    ("Cache tool results",
     "Memoize get_code / get_class lookups across leads so the same "
     "function isn\u2019t fetched multiple times in one analysis.",
     "Deferred: Current cost per-finding is already low (~$0.03\u2013$0.08). "
     "Caching adds invalidation logic. Revisit when running at scale "
     "where repeated lookups become a real cost driver."),
    ("Indexed CSV files",
     "Replace linear CSV scans with a SQLite index or in-memory hash "
     "for O(1) class/function lookups.",
     "Deferred: CSV load is <1s even for 32K-row TypeAliasLookup. "
     "Not a bottleneck today. Worth doing if CSV sizes grow 10\u00D7+."),
]
y = 2.2
for title, desc, reason in perf_items:
    tf = add_text(s, perf_left + 0.1, y, 5.4, 0.3, "", size=13,
                  color=TEXT_CLR, font_name="Segoe UI")
    add_para(tf, f"\u2022  {title}", size=14, color=TEAL, bold=True,
             font_name="Segoe UI", space_before=Pt(6))
    add_para(tf, f"   {desc}", size=11, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))
    add_para(tf, f"   \u25B8 {reason}", size=11, color=DIM,
             font_name="Segoe UI", space_before=Pt(2))
    y += 1.35

# --- Right column: Accuracy ---
acc_left = 6.8
add_section_header(s, acc_left, 1.65, 5.6, "\U0001F3AF", "Accuracy", WARN)

acc_items = [
    ("Competing / debating agents",
     "Have two LLMs independently analyze the same finding, then "
     "a judge agent reconciles disagreements.",
     "Deferred: Doubles cost per finding. Current single-agent accuracy "
     "is strong enough for triage. Add when marginal accuracy gains "
     "justify the cost increase."),
    ("Verification / auditing agents",
     "A post-synthesis reviewer that challenges the verdict, checks "
     "for logical gaps, and forces re-investigation if needed.",
     "Deferred: Adds another LLM call per finding. The synthesizer "
     "already has a confidence score and the planner\u2019s evaluate "
     "field partially serves this role. Revisit after measuring "
     "false-positive rates at scale."),
    ("Critiquing agents",
     "An adversarial agent that tries to poke holes in the evidence "
     "chain and downgrades confidence if it succeeds.",
     "Deferred: Risk of over-correction \u2014 a skeptical agent could "
     "downgrade valid findings. Needs careful calibration. Best "
     "tackled after human review data provides ground truth."),
]
y = 2.2
for title, desc, reason in acc_items:
    tf = add_text(s, acc_left + 0.1, y, 5.4, 0.3, "", size=13,
                  color=TEXT_CLR, font_name="Segoe UI")
    add_para(tf, f"\u2022  {title}", size=14, color=WARN, bold=True,
             font_name="Segoe UI", space_before=Pt(6))
    add_para(tf, f"   {desc}", size=11, color=TEXT_CLR,
             font_name="Segoe UI", space_before=Pt(2))
    add_para(tf, f"   \u25B8 {reason}", size=11, color=DIM,
             font_name="Segoe UI", space_before=Pt(2))
    y += 1.35

# Bottom principle
add_text(s, 0.5, 6.4, 12, 0.5,
         "Principle: Ship the simplest thing that works, then add complexity "
         "only when data proves it\u2019s needed.",
         size=16, color=YELLOW, bold=True, font_name="Segoe UI")


# ══════════════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════════════
prs.save(str(TARGET))
n_after = len(prs.slides)
print(f"Added {n_after - n_before} slides ({n_before} \u2192 {n_after})")
print(f"Saved \u2192 {TARGET.name}")
