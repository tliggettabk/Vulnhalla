from pptx import Presentation
prs = Presentation(r'_rearch\CID-25863-walkthrough.pptx')
for i in [6, 7]:
    slide = prs.slides[i]
    print(f'\n=== Slide {i+1} ===')
    for j, shape in enumerate(slide.shapes):
        print(f'  Shape {j}: type={shape.shape_type}, left={shape.left}, top={shape.top}, w={shape.width}, h={shape.height}')
        if shape.has_text_frame:
            tf = shape.text_frame
            for k, p in enumerate(tf.paragraphs):
                txt = p.text
                if txt.strip():
                    sz = None; bold = None; fname = None
                    if p.runs:
                        r0 = p.runs[0]
                        sz = r0.font.size
                        bold = r0.font.bold
                        fname = r0.font.name
                    print(f'    p{k}: size={sz} bold={bold} font={fname}')
                    print(f'         "{txt[:200]}"')
        elif shape.has_table:
            tbl = shape.table
            nr = len(tbl.rows); nc = len(tbl.columns)
            print(f'    [TABLE {nr}x{nc}]')
            for r in range(nr):
                cells = [tbl.cell(r, c).text[:80] for c in range(nc)]
                print(f'      row{r}: {cells}')
