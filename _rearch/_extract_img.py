from pptx import Presentation
import os
prs = Presentation(r'_rearch\CID-25863-walkthrough.pptx')
slide7 = prs.slides[6]
for j, shape in enumerate(slide7.shapes):
    if shape.shape_type == 13:  # PICTURE
        img = shape.image
        ext = img.content_type.split('/')[-1]
        out = f'_rearch/slide7_image.{ext}'
        with open(out, 'wb') as f:
            f.write(img.blob)
        print(f'Extracted image: {out} ({len(img.blob)} bytes, {img.content_type})')
        print(f'  Position: left={shape.left}, top={shape.top}, w={shape.width}, h={shape.height}')
