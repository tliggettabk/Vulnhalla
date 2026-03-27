"""Inspect both PPTX files to see slide titles side by side.

Usage:  python _rearch/inspect_pptx.py
"""
from pptx import Presentation

for label, path in [
    ("COPY (your edits)", r"_rearch\Vulnhalla_Orchestrated_Overview - Copy.pptx"),
    ("CURRENT (regenerated)", r"_rearch\Vulnhalla_Orchestrated_Overview.pptx"),
]:
    print(f"=== {label} ===")
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t and len(t) > 3:
                        print(f"  Slide {i}: {t[:90]}")
                        break
                else:
                    continue
                break
    print()
