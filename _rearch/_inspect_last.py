"""Quick inspection of the last slides in the deck."""
from pptx import Presentation

prs = Presentation(r"_rearch\Vulnhalla_Orchestrated_Overview.pptx")
print(f"Total slides: {len(prs.slides)}\n")

for idx in range(max(0, len(prs.slides) - 3), len(prs.slides)):
    slide = prs.slides[idx]
    print(f"=== Slide {idx+1} ===")
    bg = slide.background.fill
    if bg.type is not None:
        try:
            print(f"  BG: {bg.fore_color.rgb}")
        except Exception:
            print(f"  BG: (non-solid)")
    else:
        print(f"  BG: (default/none)")

    for shape in slide.shapes:
        print(f"  Shape: type={shape.shape_type}, name={shape.name}")
        if hasattr(shape, "text_frame"):
            for pi, p in enumerate(shape.text_frame.paragraphs):
                txt = p.text[:140] if p.text else ""
                font_info = ""
                if p.runs:
                    r = p.runs[0]
                    try:
                        c = r.font.color.rgb
                    except Exception:
                        c = "inherit"
                    font_info = (f"  font={r.font.name} sz={r.font.size}"
                                 f" bold={r.font.bold} color={c}")
                print(f"    P{pi}: \"{txt}\"{font_info}")
    print()
