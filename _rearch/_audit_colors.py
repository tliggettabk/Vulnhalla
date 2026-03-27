"""Audit all unique colors in the deck to find unmapped ones."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from collections import Counter

prs = Presentation(r"_rearch\Vulnhalla_Orchestrated_Overview.pptx")

fill_colors = Counter()
text_colors = Counter()
line_colors = Counter()
bg_colors = Counter()
inherit_count = 0

for idx, slide in enumerate(prs.slides):
    # Background
    try:
        bg = slide.background.fill
        if bg.type is not None:
            bg_colors[str(bg.fore_color.rgb)] += 1
    except:
        bg_colors["(error)"] += 1

    for shape in slide.shapes:
        # Fill
        try:
            f = shape.fill
            if f.type is not None:
                fill_colors[str(f.fore_color.rgb)] += 1
        except:
            pass

        # Line
        try:
            if shape.line.color and shape.line.color.rgb:
                line_colors[str(shape.line.color.rgb)] += 1
        except:
            pass

        # Text
        if hasattr(shape, "text_frame"):
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    try:
                        if run.font.color and run.font.color.rgb:
                            text_colors[str(run.font.color.rgb)] += 1
                        else:
                            inherit_count += 1
                    except:
                        inherit_count += 1

print("=== BACKGROUNDS ===")
for c, n in bg_colors.most_common():
    print(f"  {c}: {n} slides")

print("\n=== SHAPE FILLS ===")
for c, n in fill_colors.most_common():
    print(f"  {c}: {n} shapes")

print("\n=== TEXT COLORS ===")
for c, n in text_colors.most_common():
    print(f"  {c}: {n} runs")
print(f"  (inherited/no explicit color): {inherit_count} runs")

print("\n=== LINE COLORS ===")
for c, n in line_colors.most_common():
    print(f"  {c}: {n} shapes")
