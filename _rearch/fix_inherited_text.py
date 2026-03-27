"""Fix unreadable text: set all inherited/unset text colors to GitHub Dark text.

The retheme script only remapped explicitly-set RGB colors. Most text in the
deck inherits from the PowerPoint theme (defaulting to black). This script
forces every text run without an explicit color to GitHub Dark's text color.

Usage:  python _rearch/fix_inherited_text.py
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-textfix.bak.pptx"

# GitHub Dark palette
GH_TEXT   = RGBColor(0xC9, 0xD1, 0xD9)
GH_BG     = RGBColor(0x0D, 0x11, 0x17)
GH_BLUE   = RGBColor(0x58, 0xA6, 0xFF)
GH_GREEN  = RGBColor(0x3F, 0xB9, 0x50)
GH_ORANGE = RGBColor(0xD2, 0x99, 0x22)
GH_PURPLE = RGBColor(0xBC, 0x8C, 0xFF)
GH_TEAL   = RGBColor(0x39, 0xD2, 0xC0)
GH_RED    = RGBColor(0xF8, 0x51, 0x49)
GH_YELLOW = RGBColor(0xE3, 0xB3, 0x41)
GH_DIM    = RGBColor(0x8B, 0x94, 0x9E)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

# Known good colors that should NOT be touched
GOOD_COLORS = {
    (0xC9, 0xD1, 0xD9),  # GH_TEXT
    (0x0D, 0x11, 0x17),  # GH_BG (dark text on bright fills)
    (0xFF, 0xFF, 0xFF),  # WHITE
    (0x58, 0xA6, 0xFF),  # GH_BLUE
    (0x3F, 0xB9, 0x50),  # GH_GREEN
    (0xD2, 0x99, 0x22),  # GH_ORANGE
    (0xBC, 0x8C, 0xFF),  # GH_PURPLE
    (0x39, 0xD2, 0xC0),  # GH_TEAL
    (0xF8, 0x51, 0x49),  # GH_RED
    (0xE3, 0xB3, 0x41),  # GH_YELLOW
    (0x8B, 0x94, 0x9E),  # GH_DIM
}

# Stray leftover Catppuccin color → remap target
STRAY_MAP = {
    (0xAE, 0xB0, 0xBE): GH_TEXT,   # leftover Catppuccin subtext
    (0x2A, 0x2B, 0x3D): RGBColor(0x16, 0x1B, 0x22),  # leftover surface fill
}


def rgb_key(rgb):
    return (rgb[0], rgb[1], rgb[2])


shutil.copy2(TARGET, BACKUP)
print(f"Backed up -> {BACKUP.name}")

prs = Presentation(str(TARGET))
fixed_inherited = 0
fixed_stray = 0

for idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        # Fix stray fill color
        try:
            f = shape.fill
            if f.type is not None:
                key = rgb_key(f.fore_color.rgb)
                if key in STRAY_MAP:
                    f.solid()
                    f.fore_color.rgb = STRAY_MAP[key]
                    fixed_stray += 1
        except (AttributeError, TypeError):
            pass

        if not hasattr(shape, "text_frame"):
            continue

        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                has_explicit = False
                try:
                    if run.font.color and run.font.color.rgb:
                        has_explicit = True
                        key = rgb_key(run.font.color.rgb)
                        # Fix stray leftover colors
                        if key in STRAY_MAP:
                            run.font.color.rgb = STRAY_MAP[key]
                            fixed_stray += 1
                except (AttributeError, TypeError):
                    pass

                if not has_explicit:
                    run.font.color.rgb = GH_TEXT
                    fixed_inherited += 1

print(f"Fixed {fixed_inherited} inherited-color runs -> #C9D1D9")
print(f"Fixed {fixed_stray} stray leftover colors")

prs.save(str(TARGET))
print(f"Saved -> {TARGET.name}")
