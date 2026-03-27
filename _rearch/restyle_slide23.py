"""Restyle slide 23 (user's pipeline diagram) to match Catppuccin Mocha theme.

Changes colours and fonts only. Does NOT move or resize any shapes.

Usage:  python _rearch/restyle_slide23.py
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-restyle23.bak.pptx"

# ── Catppuccin Mocha palette ──
BG        = RGBColor(0x1E, 0x1E, 0x2E)
TEXT_CLR  = RGBColor(0xCD, 0xD6, 0xF4)
TITLE_CLR = RGBColor(0x89, 0xB4, 0xFA)
ACCENT    = RGBColor(0xA6, 0xE3, 0xA1)
TEAL      = RGBColor(0x94, 0xE2, 0xD5)
MAUVE     = RGBColor(0xCB, 0xA6, 0xF7)
WARN      = RGBColor(0xFA, 0xB3, 0x87)
YELLOW    = RGBColor(0xF9, 0xE2, 0xAF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG   = RGBColor(0x31, 0x32, 0x44)
SURFACE0  = RGBColor(0x31, 0x32, 0x44)
SURFACE1  = RGBColor(0x45, 0x47, 0x5A)

# Map shape names to fill + text colour
RECT_STYLES = {
    "Rectangle 1":  (TITLE_CLR, WHITE,  True),   # "Application Build Pipeline" header
    "Rectangle 4":  (SURFACE1,  TEAL,   True),   # Build
    "Rectangle 5":  (SURFACE1,  TEAL,   True),   # Run CodeQL Rules
    "Rectangle 9":  (MAUVE,     WHITE,  True),   # AI Triage (Vulnhalla)
    "Rectangle 10": (SURFACE1,  WARN,   True),   # ASM
    "Rectangle 14": (ACCENT,    BG,     True),   # Orchestrator (highlight)
    "Rectangle 15": (WARN,      WHITE,  True),   # LLM
}

shutil.copy2(TARGET, BACKUP)
print(f"Backed up -> {BACKUP.name}")

prs = Presentation(str(TARGET))
slide = prs.slides[22]  # 0-indexed → slide 23

# ── 1. Set dark background ──
bg = slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = BG
print("Set background to dark")

# ── 2. Restyle each shape ──
for shape in slide.shapes:
    name = shape.name

    # --- Rectangles ---
    if name in RECT_STYLES:
        fill_clr, txt_clr, bold_title = RECT_STYLES[name]
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_clr
        # Subtle border
        shape.line.color.rgb = fill_clr
        shape.line.width = Pt(1)

        if hasattr(shape, "text_frame"):
            for pi, p in enumerate(shape.text_frame.paragraphs):
                for run in p.runs:
                    run.font.color.rgb = txt_clr
                    run.font.name = "Segoe UI"
                    if pi == 0:
                        run.font.bold = bold_title
                # Also set paragraph-level defaults
                if p.font:
                    p.font.color.rgb = txt_clr
                    p.font.name = "Segoe UI"
        print(f"  Styled {name} -> fill={fill_clr}")

    # --- TextBoxes (labels like "Findings", "CodeQL DB") ---
    elif shape.shape_type == 17 and hasattr(shape, "text_frame"):  # TEXT_BOX
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.font.color.rgb = TEXT_CLR
                run.font.name = "Segoe UI"
            if p.font:
                p.font.color.rgb = TEXT_CLR
                p.font.name = "Segoe UI"
        print(f"  Styled {name} -> text={TEXT_CLR}")

    # --- Connectors / arrows ---
    elif shape.shape_type == 9:  # LINE
        shape.line.color.rgb = YELLOW
        shape.line.width = Pt(2)
        print(f"  Styled {name} -> line={YELLOW}")

    # --- Picture: leave as-is ---
    elif shape.shape_type == 13:
        print(f"  Skipped {name} (picture)")

prs.save(str(TARGET))
print(f"\nSaved -> {TARGET.name}")
