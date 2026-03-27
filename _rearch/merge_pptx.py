"""Insert new phase-prompt slides from Current into your edited Copy.

Usage:  python _rearch/merge_pptx.py

This script:
1. Backs up the current .pptx to .bak.pptx
2. Loads your edited Copy as the base (preserving all your changes)
3. Finds the 3 new phase-prompt slides in Current
4. Inserts them after the "3-Phase Pipeline" slide
5. Saves the result over the current .pptx
"""
import copy
import shutil
from pathlib import Path
from pptx import Presentation

REARCH  = Path(r"_rearch")
CURRENT = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
COPY    = REARCH / "Vulnhalla_Orchestrated_Overview - Copy.pptx"
BACKUP  = REARCH / "Vulnhalla_Orchestrated_Overview.bak.pptx"

# 1. Backup current
shutil.copy2(CURRENT, BACKUP)
print(f"Backed up → {BACKUP.name}")

# 2. Find the 3 new slides in Current by title keyword
curr_prs = Presentation(str(CURRENT))
new_keywords = ["Phase 1 Prompt", "Phase 2 Prompt", "Phase 3 Prompt"]
new_slide_indices = []
for i, slide in enumerate(curr_prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            if any(kw in shape.text_frame.text for kw in new_keywords):
                new_slide_indices.append(i)
                # Get title for logging
                title = ""
                for p in shape.text_frame.paragraphs:
                    if any(kw in p.text for kw in new_keywords):
                        title = p.text.strip()[:60]
                        break
                print(f"  Found new slide {i+1} in Current: {title}")
                break

if not new_slide_indices:
    print("ERROR: Could not find the 3 phase-prompt slides in Current.")
    raise SystemExit(1)

# 3. Load your edited Copy as base
target_prs = Presentation(str(COPY))
print(f"Base (your copy): {len(target_prs.slides)} slides")

# 4. Find insertion point — after "3-Phase Pipeline" slide
insert_after = 5  # default fallback
for i, slide in enumerate(target_prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            if "3-Phase Pipeline" in shape.text_frame.text:
                insert_after = i + 1
                break

print(f"Inserting after slide {insert_after}")

# 5. Insert each new slide
ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

for offset, src_idx in enumerate(new_slide_indices):
    src_slide = curr_prs.slides[src_idx]

    # Add a blank slide (will be placed at end initially)
    new_slide = target_prs.slides.add_slide(target_prs.slide_layouts[6])

    # Clear default shapes
    for shape in list(new_slide.shapes):
        new_slide.shapes._spTree.remove(shape._element)

    # Copy all shapes from source
    for shape in src_slide.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shape._element))

    # Copy background
    src_cSld = src_slide._element.find(f"{ns}cSld")
    tgt_cSld = new_slide._element.find(f"{ns}cSld")
    src_bg = src_cSld.find(f"{ns}bg") if src_cSld is not None else None
    if src_bg is not None:
        old_bg = tgt_cSld.find(f"{ns}bg")
        if old_bg is not None:
            tgt_cSld.remove(old_bg)
        tgt_cSld.insert(0, copy.deepcopy(src_bg))

    # Reorder: move new slide from end to correct position
    sldIdLst = target_prs.slides._sldIdLst
    ids = list(sldIdLst)
    last = ids[-1]
    sldIdLst.remove(last)
    target_pos = insert_after + offset
    sldIdLst.insert(target_pos, last)

    print(f"  Inserted at position {target_pos + 1}")

# 6. Save
target_prs.save(str(CURRENT))
print(f"\nSaved → {CURRENT.name}  ({len(target_prs.slides)} slides)")
print("Your manual edits from Copy are preserved.")
