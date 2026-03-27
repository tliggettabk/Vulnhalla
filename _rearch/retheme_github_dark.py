"""Retheme the entire deck from Catppuccin Mocha to GitHub Dark.

Walks every shape on every slide, remaps fill colors, text colors,
line colors. Backs up first. Does NOT move or resize anything.

Usage:  python _rearch/retheme_github_dark.py
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

REARCH = Path(r"_rearch")
TARGET = REARCH / "Vulnhalla_Orchestrated_Overview.pptx"
BACKUP = REARCH / "Vulnhalla_Orchestrated_Overview.pre-github-dark.bak.pptx"

# ── Color mapping: Catppuccin Mocha → GitHub Dark ──
# Key = old RGB tuple, Value = new RGBColor
COLOR_MAP = {
    # Backgrounds / surfaces
    (0x1E, 0x1E, 0x2E): RGBColor(0x0D, 0x11, 0x17),  # BG
    (0x31, 0x32, 0x44): RGBColor(0x16, 0x1B, 0x22),  # CODE_BG / SURFACE0
    (0x45, 0x47, 0x5A): RGBColor(0x21, 0x26, 0x2D),  # SURFACE1

    # Semantic colors
    (0x89, 0xB4, 0xFA): RGBColor(0x58, 0xA6, 0xFF),  # TITLE (blue)
    (0xCD, 0xD6, 0xF4): RGBColor(0xC9, 0xD1, 0xD9),  # TEXT (light grey)
    (0xA6, 0xE3, 0xA1): RGBColor(0x3F, 0xB9, 0x50),  # ACCENT (green)
    (0xFA, 0xB3, 0x87): RGBColor(0xD2, 0x99, 0x22),  # WARN (orange)
    (0xCB, 0xA6, 0xF7): RGBColor(0xBC, 0x8C, 0xFF),  # MAUVE (purple)
    (0xF9, 0xE2, 0xAF): RGBColor(0xE3, 0xB3, 0x41),  # YELLOW
    (0x94, 0xE2, 0xD5): RGBColor(0x39, 0xD2, 0xC0),  # TEAL
    (0xF3, 0x8B, 0xA8): RGBColor(0xF8, 0x51, 0x49),  # RED
    (0x6C, 0x70, 0x86): RGBColor(0x8B, 0x94, 0x9E),  # DIM
}

GH_BG = RGBColor(0x0D, 0x11, 0x17)


def rgb_key(rgb_color):
    """Convert RGBColor to (r, g, b) tuple for dict lookup."""
    return (rgb_color[0], rgb_color[1], rgb_color[2])


def remap(rgb_color):
    """Return the GitHub Dark equivalent, or None if no mapping."""
    if rgb_color is None:
        return None
    key = rgb_key(rgb_color)
    return COLOR_MAP.get(key)


def restyle_run(run):
    """Remap font color on a single text run."""
    try:
        if run.font.color and run.font.color.rgb:
            new = remap(run.font.color.rgb)
            if new:
                run.font.color.rgb = new
    except (AttributeError, TypeError):
        pass


def restyle_paragraph(p):
    """Remap font color on paragraph-level and all runs."""
    try:
        if p.font and p.font.color and p.font.color.rgb:
            new = remap(p.font.color.rgb)
            if new:
                p.font.color.rgb = new
    except (AttributeError, TypeError):
        pass
    for run in p.runs:
        restyle_run(run)


def restyle_text_frame(tf):
    """Remap all text colors in a text frame."""
    for p in tf.paragraphs:
        restyle_paragraph(p)


def restyle_fill(shape):
    """Remap solid fill color on a shape."""
    try:
        fill = shape.fill
        if fill.type is not None:
            try:
                rgb = fill.fore_color.rgb
                new = remap(rgb)
                if new:
                    fill.solid()
                    fill.fore_color.rgb = new
            except (AttributeError, TypeError):
                pass
    except (AttributeError, TypeError):
        pass


def restyle_line(shape):
    """Remap line/border color on a shape."""
    try:
        line = shape.line
        if line.color and line.color.rgb:
            new = remap(line.color.rgb)
            if new:
                line.color.rgb = new
    except (AttributeError, TypeError):
        pass


# ══════════════════════════════════════════════════════════════════════
# Main
# ══════════════════════════════════════════════════════════════════════
shutil.copy2(TARGET, BACKUP)
print(f"Backed up \u2192 {BACKUP.name}")

prs = Presentation(str(TARGET))
total_slides = len(prs.slides)
print(f"Slides: {total_slides}")

changes = 0

for idx, slide in enumerate(prs.slides):
    # --- Slide background ---
    bg = slide.background.fill
    try:
        if bg.type is not None and bg.fore_color.rgb:
            new_bg = remap(bg.fore_color.rgb)
            if new_bg:
                bg.solid()
                bg.fore_color.rgb = new_bg
                changes += 1
    except (AttributeError, TypeError):
        # Slide 23 has non-solid bg — force it
        bg.solid()
        bg.fore_color.rgb = GH_BG
        changes += 1

    # --- Each shape ---
    for shape in slide.shapes:
        # Fill
        restyle_fill(shape)
        changes += 1

        # Line / border
        restyle_line(shape)

        # Text
        if hasattr(shape, "text_frame"):
            restyle_text_frame(shape.text_frame)

print(f"Remapped colors across {total_slides} slides")

prs.save(str(TARGET))
print(f"Saved \u2192 {TARGET.name}")
print(f"Theme: GitHub Dark")
